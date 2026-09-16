"""Build workbook/Thai REIT Transactions 2026.xlsx from data/reits/*.json and data/analysis.json.

Sheet 1 keeps the structure of the 2025 tracker in sample/. The detail sheets carry every figure with its
extraction status (cell colour) and citation (cell comment: file:line + verbatim quote), and the
Citations sheet lists every sourced value, so each number can be rechecked against text/{FILING_ID}/*.md.
Consistency checks are live Excel formulas where the inputs are on the same row.

Usage: python pipeline/build_workbook.py
"""
from __future__ import annotations

import json
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from common import SAMPLE_WORKBOOK, WORKBOOK, utf8_console
from dataset import ANALYSIS, STATUS, cite, iter_citations, load_filings, load_json, num, val

THB_MN, THB, UNITS, PCT, SQM, YEARS, RATIO, PCT_FRAC = "#,##0.0", "#,##0.00", "#,##0", "0.00", "#,##0", "0.0", "0.00", "0.0%"
STATUS_FILL = {
    "not_found": PatternFill("solid", fgColor="E7E6E6"),
    "not_applicable": PatternFill("solid", fgColor="F2F2F2"),
    "ambiguous": PatternFill("solid", fgColor="FFE699"),
    "derived": PatternFill("solid", fgColor="DDEBF7"),
}
RESULT_FILL = {"pass": "C6EFCE", "warn": "FFEB9C", "fail": "FFC7CE"}
TOP = Alignment(vertical="top", wrap_text=True)
SECTOR_LABEL = {"Industrial & Logistics": "Industrial", "Residential & Serviced Apartment": "Residence"}
ROLES = [
    ("Sponsor", "sponsors"), ("REIT manager", "reit_manager"), ("Trustee", "trustee"),
    ("Property manager", "property_managers"), ("Financial advisor", "financial_advisors"),
    ("Lead underwriter", "lead_underwriters"), ("Underwriter", "underwriters"), ("Appraiser", "appraisers"),
    ("Auditor", "auditor"), ("Legal advisor", "legal_advisors"), ("Lender", "lenders"),
    ("Seller / lessor", "sellers_or_lessors"), ("Operator / master lessee", "operators_or_master_lessees"),
]


class SV:
    """A value from the extraction with its status, citation and note."""

    def __init__(self, obj):
        is_obj = isinstance(obj, dict)
        self.value = val(obj) if is_obj else obj
        self.status = obj.get("status") if is_obj else None
        self.source = obj.get("source") if is_obj else None
        self.note = obj.get("note", "") if is_obj else ""


class F(str):
    """An Excel formula."""


def clean(x):
    if isinstance(x, str):
        x = ILLEGAL_CHARACTERS_RE.sub("", x)
        return " " + x if x.startswith("=") and not isinstance(x, F) else x
    if isinstance(x, (list, dict)):
        return json.dumps(x, ensure_ascii=False)
    return x


def put(ws, row: int, col: int, item, fmt: str | None = None) -> None:
    cell = ws.cell(row=row, column=col)
    if isinstance(item, SV):
        cell.value = clean(item.value)
        if item.status in STATUS_FILL:
            cell.fill = STATUS_FILL[item.status]
        lines = [item.status or ""]
        if item.source:
            lines.append(cite(item.source))
            if item.source.get("quote"):
                lines.append("“" + item.source["quote"][:240] + "”")
        if item.note:
            lines.append(item.note[:240])
        if item.source or item.note:
            comment = Comment(ILLEGAL_CHARACTERS_RE.sub("", "\n".join(x for x in lines if x)), "extraction")
            comment.width, comment.height = 360, 140
            cell.comment = comment
    else:
        cell.value = clean(item)
    if fmt and (isinstance(cell.value, (int, float)) or isinstance(item, F)):
        cell.number_format = fmt
    cell.alignment = TOP


def letter(columns, header: str) -> str:
    return get_column_letter([c[0] for c in columns].index(header) + 1)


def add_sheet(wb, title: str, table: str, columns: list[tuple], rows: list[list], note: str | None = None):
    """columns: (header, width, number_format). A row item may be a callable(row_number) returning a formula."""
    ws = wb.create_sheet(title)
    header_row = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="595959")
        header_row = 3
    for i, (header, width, _) in enumerate(columns, 1):
        cell = ws.cell(header_row, i, header)
        cell.font = Font(bold=True)
        cell.alignment = TOP
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, row in enumerate(rows, header_row + 1):
        for c, ((_, _, fmt), item) in enumerate(zip(columns, row), 1):
            put(ws, r, c, item(r) if callable(item) else item, fmt)
    if rows:
        t = Table(displayName=table, ref=f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = ws.cell(header_row + 1, 3)
    return ws


def d_mon(iso: str | None) -> str:
    try:
        return datetime.strptime(iso, "%Y-%m-%d").strftime("%-d %b %Y")
    except (TypeError, ValueError):
        try:
            return datetime.strptime(iso, "%Y-%m-%d").strftime("%#d %b %Y")
        except (TypeError, ValueError):
            return iso or ""


def join_names(items) -> str:
    if isinstance(items, dict):
        items = [items]
    return "; ".join(p["name_en"] for p in items or [] if isinstance(p, dict))


# ---------------------------------------------------------------- sheet 1: tracker (2025 structure)

def tracker(wb, filings, analysis, status):
    ws = wb.active
    ws.title = "reit-transaction-2026"
    widths = {"A": 36.3, "B": 11, "C": 12, "D": 16.7, "E": 11, "F": 15.7, "G": 42, "H": 14, "I": 16, "J": 34, "K": 34, "L": 22}
    if SAMPLE_WORKBOOK.exists():
        sample = load_workbook(SAMPLE_WORKBOOK).active
        widths.update({k: v.width for k, v in sample.column_dimensions.items() if v.width and k in widths})
    ws["A1"] = "Thai REIT Transactions 2026"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Built {datetime.now():%d %b %Y %H:%M} from the text filings in filings/ only. Hover a value for its citation; detail sheets follow."
    ws["A2"].font = Font(italic=True, color="595959")
    headers = ["REIT Name", "Symbol", "Industry", "Transaction Type", "Seller", "Status", "Asset Details / Project Name",
               "Additional Asset Acquisition No.", "Transaction Size / Value in THB mn", "Funding Method", "Key Dates (2026)", "Filing ID"]
    for i, h in enumerate(headers, 1):
        ws.cell(3, i, h).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[get_column_letter(i)]
    row = 3
    for f in filings:
        fid = f["filing"]["filing_id"]
        a = analysis["filings"].get(fid, {})
        m = a.get("metrics", {})
        sec = (status.get("filings") or {}).get(fid) or {}
        sellers = f["parties"]["sellers_or_lessors"]
        if not sellers:
            seller = "n/a"
        elif any(p.get("related_to_sponsor") is True for p in sellers):
            seller = "Sponsor"
        elif all(p.get("related_to_sponsor") is False for p in sellers):
            seller = "3rd Parties"
        else:
            seller = "Not stated"
        stage = sec.get("stage")
        status_txt = {"closed": "Offering closed", "withdrawn": "Withdrawn"}.get(stage, "In progress")
        new = f["new_assets"]
        details = f"{f['sector']['primary']}: {len(new)} asset{'s' if len(new) != 1 else ''}:\n" + "\n".join(
            f"{i}. {x['name_en']} ({val(x['investment_type']) or 'tenure n/a'})" for i, x in enumerate(new, 1))
        is_ipo = f["filing"]["offering_type"] == "IPO"
        add_no = 0 if is_ipo else (num(f["filing"].get("additional_investment_no")) or "n/a")
        units, pmax, debt = m.get("units_offered_max"), m.get("price_max_thb"), m.get("new_debt_thb_mn")
        bits = [b for b in (f"up to {units / 1e6:,.1f}M units" if units else "", f"at up to THB {pmax:,.2f}/unit" if pmax else "") if b]
        funding = ("Initial Public Offering (IPO)" if is_ipo else "Capital Increase (PO)") + (f" ({' '.join(bits)})" if bits else "")
        if debt:
            funding += f" / Debt up to THB {debt:,.0f}M"
        if any(k in (m.get("funding_thb_mn") or {}) for k in ("internal_cash", "security_deposits")):
            funding += " / Internal cash & deposits"
        dates = [f"Filing Date: {d_mon(sec.get('first_filed'))}" if sec.get("first_filed") else ""]
        if sec.get("last_amended"):
            dates.append(f"Last amended: {d_mon(sec['last_amended'])}")
        if sec.get("effective"):
            dates.append(f"Effective: {d_mon(sec['effective'])}")
        if sec.get("offer_start"):
            dates.append(f"Offer period: {d_mon(sec['offer_start'])} – {d_mon(sec.get('offer_end'))}")
        dates.append(f"SEC stage ({d_mon(status.get('as_of'))}): {sec.get('stage_label', 'not checked')}")
        row += 1
        values = [SV(f["filing"]["reit_name_en"]), f["filing"]["ticker"], SECTOR_LABEL.get(f["sector"]["primary"], f["sector"]["primary"]),
                  f["filing"]["offering_type"], seller, status_txt, details, add_no, SV(f["offering"]["total_investment_value_thb_mn"]),
                  funding, "\n".join(x for x in dates if x), fid]
        for c, v in enumerate(values, 1):
            put(ws, row, c, v, "#,##0" if c == 9 else None)
    if row > 3:
        t = Table(displayName="Table1", ref=f"A3:L{row}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
        ws.cell(row + 2, 8, "Total")
        ws.cell(row + 2, 9, "=SUBTOTAL(9,Table1[Transaction Size / Value in THB mn])").number_format = "#,##0.00"
        ws.cell(row + 3, 8, "Count")
        ws.cell(row + 3, 9, "=SUBTOTAL(2,Table1[Transaction Size / Value in THB mn])").number_format = "#,##0"


# ---------------------------------------------------------------- detail sheets

def build(filings, analysis, status) -> Workbook:
    wb = Workbook()
    tracker(wb, filings, analysis, status)
    tick = lambda f: f["filing"]["ticker"]  # noqa: E731
    fid = lambda f: f["filing"]["filing_id"]  # noqa: E731

    # Offering & capital
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Type", 6, None), ("Max units", 14, UNITS), ("Min price (THB)", 10, THB),
            ("Max price (THB)", 10, THB), ("Final price (THB)", 10, THB), ("Par (THB)", 10, "0.0000"), ("Offering size (THB mn)", 13, THB_MN),
            ("Check: units x price (THB mn)", 14, THB_MN), ("Check: diff vs stated", 10, PCT_FRAC), ("Total investment (THB mn)", 13, THB_MN),
            ("New units funding (THB mn)", 13, THB_MN), ("New debt (THB mn)", 12, THB_MN), ("Cash / deposits / other (THB mn)", 13, THB_MN),
            ("Funding total (THB mn)", 13, THB_MN), ("Est. transaction costs (THB mn)", 13, THB_MN), ("Projected DPU (THB)", 11, "0.0000"),
            ("Yield basis price (THB)", 11, THB), ("Projected yield (%)", 10, PCT), ("Check: DPU / price (%)", 11, PCT), ("Projection period", 26, None),
            ("LTV before (%)", 10, PCT), ("LTV after (%)", 10, PCT), ("LTV limit (%)", 10, PCT), ("Units before", 14, UNITS),
            ("Units after (max)", 14, UNITS), ("NAV per unit (THB)", 11, "0.0000"), ("Investor eligibility", 34, None),
            ("Subscription ratio", 26, None), ("Sponsor commitment", 34, None), ("Loan terms", 34, None)]
    L = lambda h: letter(cols, h)  # noqa: E731
    rows = []
    for f in filings:
        off, cap = f["offering"], f["capital_structure"]
        by_type = {}
        for fs in off["funding_sources"]:
            by_type.setdefault(fs["type"], []).append(fs["amount_thb_mn"])

        def funding_cell(types):
            objs = [o for t in types for o in by_type.get(t, [])]
            if len(objs) == 1:
                return SV(objs[0])
            nums = [num(o) for o in objs if num(o) is not None]
            return sum(nums) if nums else None

        price = off["offering_price_per_unit"]
        rows.append([
            fid(f), tick(f), f["filing"]["offering_type"], SV(off["units_offered_max"]), SV(price["min_thb"]), SV(price["max_thb"]),
            SV(price["final_thb"]), SV(off["par_value_per_unit_thb"]), SV(off["offering_size_thb_mn"]),
            lambda r: F(f'=IF(AND(ISNUMBER({L("Max units")}{r}),OR(ISNUMBER({L("Final price (THB)")}{r}),ISNUMBER({L("Max price (THB)")}{r}))),'
                        f'{L("Max units")}{r}*IF(ISNUMBER({L("Final price (THB)")}{r}),{L("Final price (THB)")}{r},{L("Max price (THB)")}{r})/1000000,"")'),
            lambda r: F(f'=IF(AND(ISNUMBER({L("Check: units x price (THB mn)")}{r}),ISNUMBER({L("Offering size (THB mn)")}{r})),'
                        f'{L("Check: units x price (THB mn)")}{r}/{L("Offering size (THB mn)")}{r}-1,"")'),
            SV(off["total_investment_value_thb_mn"]), funding_cell(["new_units"]), funding_cell(["debt"]),
            funding_cell(["internal_cash", "security_deposits", "other"]),
            lambda r: F(f'=SUM({L("New units funding (THB mn)")}{r}:{L("Cash / deposits / other (THB mn)")}{r})'),
            SV(off["estimated_transaction_costs_thb_mn"]), SV(off["projected_distribution"]["dpu_thb"]),
            SV(off["projected_distribution"]["basis_price_thb"]), SV(off["projected_distribution"]["distribution_yield_pct"]),
            lambda r: F(f'=IF(AND(ISNUMBER({L("Projected DPU (THB)")}{r}),ISNUMBER({L("Yield basis price (THB)")}{r})),'
                        f'{L("Projected DPU (THB)")}{r}/{L("Yield basis price (THB)")}{r}*100,"")'),
            SV(off["projected_distribution"]["projection_period"]), SV(cap["ltv_before_pct"]), SV(cap["ltv_after_pct"]), SV(cap["ltv_limit_pct"]),
            SV(cap["units_outstanding_before"]), SV(cap["units_outstanding_after_max"]), SV(cap["nav_per_unit_thb"]),
            SV(f["filing"]["investor_eligibility"]), SV(off.get("subscription_ratio")), SV(off["sponsor_commitment"]), SV(cap["loan_terms"]),
        ])
    add_sheet(wb, "Offering & Capital", "tblOffering", cols, rows,
              "Grey = not found in text filings · yellow = ambiguous / draft placeholder · blue = derived. Check columns are live formulas.")

    # Use of proceeds & allocation & timeline
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Kind", 16, None), ("Item", 48, None), ("Amount (THB mn)", 14, THB_MN),
            ("Units", 14, UNITS), ("% of offer", 10, PCT), ("Date", 16, None), ("Note", 40, None), ("Source", 40, None)]
    rows = []
    for f in filings:
        off = f["offering"]
        for u in off["use_of_proceeds"]:
            rows.append([fid(f), tick(f), "Use of proceeds", u["item"], u["amount_thb_mn"], None, None, None, u.get("note", ""), cite(u["source"])])
        for a in off["allocation"]:
            rows.append([fid(f), tick(f), "Allocation", a["tranche"], None, a.get("units"), a.get("pct_of_offer"), None, a.get("note", ""), cite(a["source"])])
        for t in off["timeline"]:
            rows.append([fid(f), tick(f), "Timeline", t["event"], None, None, None, t.get("date"), "indicative" if t.get("indicative") else "", cite(t["source"])])
        for fs in off["funding_sources"]:
            rows.append([fid(f), tick(f), "Funding source", fs["type"], SV(fs["amount_thb_mn"]), None, None, None, fs.get("note", ""), cite(fs["amount_thb_mn"].get("source"))])
    add_sheet(wb, "Proceeds & Timeline", "tblProceeds", cols, rows)

    # Parties
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Role", 20, None), ("Name", 44, None), ("Name (TH)", 44, None),
            ("Related to sponsor", 10, None), ("Detail", 40, None), ("Source", 40, None)]
    rows = []
    for f in filings:
        for label, key in ROLES:
            items = f["parties"][key]
            for p in ([items] if isinstance(items, dict) else items or []):
                rows.append([fid(f), tick(f), label, p["name_en"], p.get("name_th"), p.get("related_to_sponsor"), p.get("role_detail"), cite(p["source"])])
        for role in f["parties"]["not_found_roles"]:
            rows.append([fid(f), tick(f), role, "(not named in text filings)", None, None, None, ""])
    add_sheet(wb, "Parties", "tblParties", cols, rows)

    # New assets
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Asset ID", 16, None), ("Asset", 36, None), ("Asset (TH)", 36, None), ("Type", 18, None),
            ("Address", 36, None), ("District", 16, None), ("Province", 16, None), ("Zone", 16, None), ("Investment type", 18, None),
            ("Tenure (years)", 10, YEARS), ("Lease start", 12, None), ("Lease expiry", 12, None), ("Components", 44, None), ("Land (sqm)", 12, SQM),
            ("Land as filed", 16, None), ("Gross area (sqm)", 12, SQM), ("Leasable area (sqm)", 12, SQM), ("Rooms", 8, UNITS), ("Buildings", 8, UNITS),
            ("Completion / age", 18, None), ("Seller / lessor", 30, None), ("Acquisition price (THB mn)", 13, THB_MN),
            ("Appraiser 1", 28, None), ("Appraisal 1 (THB mn)", 13, THB_MN), ("Method 1", 22, None), ("Valuation date 1", 12, None),
            ("Appraiser 2", 28, None), ("Appraisal 2 (THB mn)", 13, THB_MN), ("Method 2", 22, None), ("Valuation date 2", 12, None),
            ("Check: price vs lower appraisal", 12, PCT_FRAC), ("Latest occupancy (%)", 10, PCT), ("Occupancy as of", 12, None),
            ("WALE (years)", 9, "0.00"), ("Lease structure", 40, None), ("Rental guarantee", 30, None), ("Key tenants", 40, None),
            ("Tenant concentration", 30, None), ("Lease expiry profile", 36, None), ("Encumbrances", 30, None), ("Notes", 40, None)]
    L = lambda h: letter(cols, h)  # noqa: E731
    rows = []
    for f in filings:
        for a in f["new_assets"]:
            ap = a["appraisals"] + [{}, {}]
            loc = a["location"]
            occ = sorted((o for o in a["occupancy"] if o.get("occupancy_pct") is not None), key=lambda o: o["as_of"])
            appraisal = lambda i, k: SV({"value": ap[i].get(k), "status": "found", "source": ap[i]["source"]}) if ap[i].get("source") else ap[i].get(k)  # noqa: E731
            rows.append([
                fid(f), tick(f), a["asset_id"], a["name_en"], a.get("name_th"), a["asset_type"], SV({"value": loc.get("address"), "status": "found", "source": loc.get("source")}) if loc.get("source") else loc.get("address"),
                loc.get("district"), loc.get("province"), loc.get("zone"), SV(a["investment_type"]), SV(a["tenure_years"]), SV(a.get("lease_start")),
                SV(a["lease_expiry"]), "\n".join(f"{c['component']}: {c['right']}" + (f", {c['tenure_years']} yrs" if c.get("tenure_years") else "")
                                                  + (f", {c.get('start_date') or '?'} → {c.get('expiry_date') or '?'}" if c.get("expiry_date") or c.get("start_date") else "")
                                                  for c in a["components"]),
                SV({"value": a["land_area"].get("value_sqm"), **{k: a["land_area"].get(k) for k in ("status", "source", "note")}}), a["land_area"].get("as_filed"),
                SV(a["building_area"]["gross_sqm"]), SV(a["building_area"]["leasable_sqm"]), SV(a["building_area"]["rooms"]), SV(a["building_area"]["buildings"]),
                SV(a.get("completion_or_age")), SV(a["seller_or_lessor"]), SV(a["acquisition_price_thb_mn"]),
                ap[0].get("appraiser"), appraisal(0, "value_thb_mn"), ap[0].get("method"), ap[0].get("valuation_date"),
                ap[1].get("appraiser"), appraisal(1, "value_thb_mn"), ap[1].get("method"), ap[1].get("valuation_date"),
                lambda r: F(f'=IF(AND(ISNUMBER({L("Acquisition price (THB mn)")}{r}),COUNT({L("Appraisal 1 (THB mn)")}{r},{L("Appraisal 2 (THB mn)")}{r})>0),'
                            f'{L("Acquisition price (THB mn)")}{r}/MIN({L("Appraisal 1 (THB mn)")}{r},{L("Appraisal 2 (THB mn)")}{r})-1,"")'),
                occ[-1]["occupancy_pct"] if occ else None, occ[-1]["as_of"] if occ else None, SV(a["wale_years"]), SV(a["lease_structure"]),
                SV(a.get("rental_guarantee")), "\n".join(t["name"] + (f" ({t['pct_of_area']}% area)" if t.get("pct_of_area") else "")
                                                         + (f" ({t['pct_of_revenue']}% revenue)" if t.get("pct_of_revenue") else "") for t in a["key_tenants"]),
                SV(a.get("tenant_concentration")), "; ".join(f"{e['period']}: {e['pct']}% ({e['basis']})" for e in a["lease_expiry_profile"]),
                SV(a.get("encumbrances")), "\n".join(a.get("notes", [])),
            ])
    add_sheet(wb, "New Assets", "tblNewAssets", cols, rows,
              "One row per asset to be acquired. Appraisals beyond two, tenants and performance are on their own sheets.")

    # Appraisals (all)
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Asset ID", 16, None), ("Asset", 34, None), ("Appraiser", 34, None),
            ("Valuation date", 12, None), ("Method", 26, None), ("Value (THB mn)", 13, THB_MN), ("Discount rate (%)", 10, PCT),
            ("Terminal cap rate (%)", 10, PCT), ("Key assumptions", 50, None), ("Acquisition price (THB mn)", 13, THB_MN),
            ("Check: price vs this appraisal", 12, PCT_FRAC), ("Note", 30, None), ("Source", 40, None)]
    L = lambda h: letter(cols, h)  # noqa: E731
    rows = []
    for f in filings:
        for a in f["new_assets"]:
            for ap in a["appraisals"]:
                rows.append([fid(f), tick(f), a["asset_id"], a["name_en"], ap["appraiser"], ap.get("valuation_date"), ap.get("method"),
                             ap.get("value_thb_mn"), ap.get("discount_rate_pct"), ap.get("terminal_cap_rate_pct"), ap.get("key_assumptions"),
                             SV(a["acquisition_price_thb_mn"]),
                             lambda r: F(f'=IF(AND(ISNUMBER({L("Acquisition price (THB mn)")}{r}),ISNUMBER({L("Value (THB mn)")}{r})),'
                                         f'{L("Acquisition price (THB mn)")}{r}/{L("Value (THB mn)")}{r}-1,"")'),
                             ap.get("note", ""), cite(ap["source"])])
    add_sheet(wb, "Appraisals", "tblAppraisals", cols, rows)

    # Asset performance
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Asset ID", 16, None), ("Asset", 34, None), ("Period", 11, None), ("Projection", 9, None),
            ("Revenue (THB mn)", 12, THB_MN), ("NOI (THB mn)", 12, THB_MN), ("NOI definition", 26, None), ("NOI margin", 9, PCT_FRAC),
            ("Occupancy (%)", 10, PCT), ("ADR (THB)", 10, UNITS), ("RevPAR (THB)", 10, UNITS), ("Avg rent (THB/sqm/month)", 12, THB),
            ("Other metrics", 34, None), ("Source", 40, None)]
    L = lambda h: letter(cols, h)  # noqa: E731
    rows = []
    for f in filings:
        for a in f["new_assets"]:
            for p in a["historical_performance"]:
                rows.append([fid(f), tick(f), a["asset_id"], a["name_en"], p["period"], "yes" if p["is_projection"] else "", p.get("revenue_thb_mn"),
                             p.get("noi_thb_mn"), p.get("noi_definition"),
                             lambda r: F(f'=IF(AND(ISNUMBER({L("NOI (THB mn)")}{r}),ISNUMBER({L("Revenue (THB mn)")}{r}),{L("Revenue (THB mn)")}{r}<>0),'
                                         f'{L("NOI (THB mn)")}{r}/{L("Revenue (THB mn)")}{r},"")'),
                             p.get("occupancy_pct"), p.get("adr_thb"), p.get("revpar_thb"), p.get("avg_rent_thb_sqm_month"), p.get("other_metrics"), cite(p["source"])])
    add_sheet(wb, "Asset Performance", "tblAssetPerf", cols, rows)

    # Tenants
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Asset ID", 16, None), ("Tenant", 36, None), ("Tenant (TH)", 30, None), ("Business", 26, None),
            ("Area (sqm)", 12, SQM), ("% of area", 9, PCT), ("% of revenue", 9, PCT), ("Lease expiry", 12, None), ("Related to sponsor", 10, None), ("Source", 40, None)]
    rows = [[fid(f), tick(f), a["asset_id"], t["name"], t.get("name_th"), t.get("business"), t.get("area_sqm"), t.get("pct_of_area"), t.get("pct_of_revenue"),
             t.get("lease_expiry"), t.get("related_to_sponsor"), cite(t["source"])]
            for f in filings for a in f["new_assets"] for t in a["key_tenants"]]
    add_sheet(wb, "Tenants", "tblTenants", cols, rows)

    # Existing portfolio (PO)
    po = [f for f in filings if f.get("existing_portfolio")]
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("As of", 12, None), ("Listing date", 12, None), ("No. of assets", 9, UNITS),
            ("Total assets (THB mn)", 13, THB_MN), ("Investment value (THB mn)", 13, THB_MN), ("NAV (THB mn)", 13, THB_MN),
            ("NAV per unit (THB)", 11, "0.0000"), ("Market price (THB)", 11, THB), ("Market price as of", 12, None), ("Check: P/NAV (x)", 9, RATIO),
            ("Latest appraisal total (THB mn)", 13, THB_MN), ("Max PO price (THB)", 11, THB), ("Check: PO price vs market", 11, PCT_FRAC),
            ("Deal size (THB mn)", 13, THB_MN), ("Check: deal / total assets", 11, PCT_FRAC), ("Trailing FY DPU (THB)", 11, "0.0000"),
            ("Check: trailing yield at market (%)", 12, PCT), ("Notes", 40, None)]
    L = lambda h: letter(cols, h)  # noqa: E731
    rows = []
    for f in po:
        ep = f["existing_portfolio"]
        ex = analysis["filings"].get(fid(f), {}).get("metrics", {}).get("existing", {})
        rows.append([
            fid(f), tick(f), SV(ep["as_of"]), SV(ep.get("listing_date")), SV(ep["number_of_assets"]), SV(ep["total_asset_value_thb_mn"]),
            SV(ep.get("investment_value_thb_mn")), SV(ep.get("nav_thb_mn")), SV(ep["nav_per_unit_thb"]), SV(ep["market_price_per_unit_thb"]),
            ep["market_price_per_unit_thb"].get("as_of"),
            lambda r: F(f'=IF(AND(ISNUMBER({L("Market price (THB)")}{r}),ISNUMBER({L("NAV per unit (THB)")}{r})),{L("Market price (THB)")}{r}/{L("NAV per unit (THB)")}{r},"")'),
            SV(ep["latest_appraisal_total_thb_mn"]), SV(f["offering"]["offering_price_per_unit"]["max_thb"]),
            lambda r: F(f'=IF(AND(ISNUMBER({L("Max PO price (THB)")}{r}),ISNUMBER({L("Market price (THB)")}{r})),{L("Max PO price (THB)")}{r}/{L("Market price (THB)")}{r}-1,"")'),
            SV(f["offering"]["total_investment_value_thb_mn"]),
            lambda r: F(f'=IF(AND(ISNUMBER({L("Deal size (THB mn)")}{r}),ISNUMBER({L("Total assets (THB mn)")}{r})),{L("Deal size (THB mn)")}{r}/{L("Total assets (THB mn)")}{r},"")'),
            ex.get("trailing_dpu_thb"),
            lambda r: F(f'=IF(AND(ISNUMBER({L("Trailing FY DPU (THB)")}{r}),ISNUMBER({L("Market price (THB)")}{r})),{L("Trailing FY DPU (THB)")}{r}/{L("Market price (THB)")}{r}*100,"")'),
            "\n".join(ep.get("notes", [])),
        ])
    add_sheet(wb, "Existing Portfolio", "tblPortfolio", cols, rows, "PO filings only; figures exclude the assets being acquired.")

    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Asset ID", 16, None), ("Asset", 36, None), ("Asset (TH)", 30, None), ("Type", 18, None),
            ("Location", 30, None), ("Province", 16, None), ("Investment type", 18, None), ("Lease expiry", 12, None), ("Acquired", 22, None),
            ("Land (sqm)", 12, SQM), ("Leasable area (sqm)", 12, SQM), ("Rooms", 8, UNITS), ("Acquisition price (THB mn)", 13, THB_MN),
            ("Latest appraisal (THB mn)", 13, THB_MN), ("Appraisal date", 12, None), ("Appraiser", 26, None), ("Occupancy (%)", 10, PCT),
            ("Occupancy as of", 12, None), ("Key tenants", 34, None), ("Notes", 34, None), ("Source", 40, None)]
    rows = [[fid(f), tick(f), x["asset_id"], x["name_en"], x.get("name_th"), x["asset_type"], x.get("location"), x.get("province"), x.get("investment_type"),
             x.get("lease_expiry"), x.get("acquired"), x.get("land_area_sqm"), x.get("leasable_area_sqm"), x.get("rooms"), x.get("acquisition_price_thb_mn"),
             x.get("latest_appraisal_thb_mn"), x.get("appraisal_date"), x.get("appraiser"), x.get("occupancy_pct"), x.get("occupancy_as_of"),
             x.get("key_tenants"), x.get("notes"), "; ".join([cite(x["source"])] + [cite(s) for s in x.get("extra_sources", [])])]
            for f in po for x in f["existing_portfolio"]["assets"]]
    add_sheet(wb, "Existing Assets", "tblExistingAssets", cols, rows)

    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Series", 12, None), ("Period", 11, None), ("DPU (THB)", 10, "0.0000"), ("DPU component", 16, None),
            ("Total revenue (THB mn)", 13, THB_MN), ("Rental & service revenue (THB mn)", 13, THB_MN), ("Net investment income (THB mn)", 13, THB_MN),
            ("Net increase in net assets (THB mn)", 13, THB_MN), ("Occupancy (%)", 10, PCT), ("Other metrics", 34, None), ("Source", 40, None)]
    rows = []
    for f in po:
        ep = f["existing_portfolio"]
        rows += [[fid(f), tick(f), "Distribution", d["period"], d["dpu_thb"], d.get("component"), None, None, None, None, None, None, cite(d["source"])]
                 for d in ep["historical_dpu"]]
        rows += [[fid(f), tick(f), "Operations", p["period"], None, None, p.get("total_revenue_thb_mn"), p.get("rental_and_service_revenue_thb_mn"),
                  p.get("net_investment_income_thb_mn"), p.get("net_increase_in_net_assets_thb_mn"), p.get("occupancy_pct"), p.get("other_metrics"), cite(p["source"])]
                 for p in ep["historical_performance"]]
        rows += [[fid(f), tick(f), "Occupancy", o["as_of"], None, None, None, None, None, None, o.get("occupancy_pct"), o.get("basis"), cite(o["source"])]
                 for o in ep.get("portfolio_occupancy", [])]
    add_sheet(wb, "Portfolio History", "tblPortfolioHistory", cols, rows)

    # Risks & notes
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Category", 20, None), ("Risk", 40, None), ("Summary", 70, None),
            ("Asset-specific", 9, None), ("Applies to", 20, None), ("Source", 40, None)]
    rows = [[fid(f), tick(f), k["category"], k["title"], k["summary"], k["asset_specific"], ", ".join(k.get("applies_to", [])), cite(k["source"])]
            for f in filings for k in f["risks"]]
    add_sheet(wb, "Risks", "tblRisks", cols, rows)
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Topic", 24, None), ("Observation", 80, None), ("Source", 40, None)]
    rows = [[fid(f), tick(f), n["topic"], n["observation"], cite(n.get("source"))] for f in filings for n in f["analyst_notes"]]
    rows += [["(all)", "", f"Insight: {i['kind']}", i["text"], "data/analysis.json"] for i in analysis.get("insights", [])]
    add_sheet(wb, "Analyst Notes", "tblNotes", cols, rows)

    # Validation
    cols = [("Filing ID", 22, None), ("Symbol", 10, None), ("Source of check", 18, None), ("Check / field", 40, None), ("Result", 11, None), ("Detail", 90, None)]
    rows = []
    for f in filings:
        a = analysis["filings"].get(fid(f), {})
        rows += [[fid(f), tick(f), "consolidate.py", c["label"], c["result"], c["detail"]] for c in a.get("checks", [])]
        rows += [[fid(f), tick(f), "extraction agent", c["check"], c["result"], c.get("detail", "")] for c in f["extraction"]["self_checks"]]
        rows += [[fid(f), tick(f), "not found", n["field"], "not found", n["reason"]] for n in f["extraction"]["not_found"]]
        rows += [[fid(f), tick(f), "ambiguity", n["field"], "ambiguous", n["detail"]] for n in f["extraction"]["ambiguities"]]
        rows += [[fid(f), tick(f), "draft revision", r["field"], "revised", f"{r['earlier_value']} ({r['earlier_file']}) → {r['latest_value']} ({r['latest_file']}) {r.get('note', '')}"]
                 for r in f["filing"]["revision_notes"]]
    ws = add_sheet(wb, "Validation", "tblValidation", cols, rows)
    for r in range(2, len(rows) + 2):
        colour = RESULT_FILL.get(ws.cell(r, 5).value)
        if colour:
            ws.cell(r, 5).fill = PatternFill("solid", fgColor=colour)

    # SEC status
    cols = [("Filing ID", 22, None), ("Stage", 26, None), ("SEC offering type", 10, None), ("First filed", 12, None), ("Last amended", 12, None),
            ("Effective", 12, None), ("Offer start", 12, None), ("Offer end", 12, None), ("Remark", 30, None), ("Issuer (as listed)", 60, None), ("SEC page", 60, None)]
    rows = [[k, s.get("stage_label"), s.get("offering_type"), s.get("first_filed"), s.get("last_amended"), s.get("effective"), s.get("offer_start"),
             s.get("offer_end"), s.get("remark"), s.get("issuer"), s.get("detail_url")] for k, s in (status.get("filings") or {}).items()]
    add_sheet(wb, "SEC Status", "tblSecStatus", cols, rows,
              f"Source: {status.get('source_url', 'not fetched')} · fetched {status.get('fetched_at', 'n/a')}")

    # Citations
    cols = [("Filing ID", 22, None), ("Field path", 56, None), ("Value", 40, None), ("Unit", 10, None), ("Status", 12, None), ("File", 44, None),
            ("Line", 7, None), ("Section", 36, None), ("Quote", 70, None), ("Note", 40, None)]
    rows = []
    for f in filings:
        for path, value, unit, st, src, note in iter_citations(f):
            if path.startswith("$.extraction"):
                continue
            src = src or {}
            rows.append([fid(f), path, value, unit, st, src.get("file"), src.get("line"), src.get("section"), src.get("quote"), note])
    add_sheet(wb, "Citations", "tblCitations", cols, rows, "Every sourced value and cited row. Files are the Markdown reading copies in text/{Filing ID}/.")

    # About
    ws = wb.create_sheet("About")
    lines = [
        ("Thai REIT Transactions 2026 — how to read this workbook", True),
        (f"Built {datetime.now():%Y-%m-%d %H:%M} by pipeline/build_workbook.py from data/reits/*.json (schema 1.0.0) and data/analysis.json.", False),
        ("Only the text (Word) filings in filings/ were used; PDFs (appraisal reports, audited statements, trust deeds) were not read.", False),
        ("Units: THB mn = million baht; per-unit figures in THB; percentages as 0–100 unless the column shows %; areas in sqm; dates ISO (Gregorian).", False),
        ("Cell colours: grey = not found in text filings · light grey = not applicable · yellow = ambiguous or draft placeholder · blue = derived by the extractor.", False),
        ("Comments on value cells give status, citation (file:line) and the verbatim Thai quote; the Citations sheet lists all of them.", False),
        ("'Check:' columns are live formulas; the Validation sheet lists the pipeline's cross-checks, the extractor's self-checks and every field not found.", False),
        ("To add filings: drop the SEC folder (ALL_<no>_Vxx) into inbox/ and run  python pipeline/run.py  — see README.md.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        ws.cell(i, 1, text).font = Font(bold=bold, size=13 if bold else 11)
    ws.column_dimensions["A"].width = 150
    return wb


def main() -> None:
    utf8_console()
    analysis = load_json(ANALYSIS, None)
    if analysis is None:
        raise SystemExit("data/analysis.json missing: run pipeline/consolidate.py first")
    status = load_json(STATUS, {})
    by_id = {f["filing"]["filing_id"]: f for f in load_filings()}
    ordered = [by_id[i] for i in analysis["order"] if i in by_id] + [f for i, f in by_id.items() if i not in analysis["order"]]
    wb = build(ordered, analysis, status)
    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(WORKBOOK)
    except PermissionError:
        raise SystemExit(f"Close '{WORKBOOK.name}' in Excel and run again.")
    print(f"{len(ordered)} filings -> {WORKBOOK.relative_to(WORKBOOK.parents[1])}")


if __name__ == "__main__":
    main()
